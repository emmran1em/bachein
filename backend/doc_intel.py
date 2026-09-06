"""Bachein Document Intelligence — extraction (PyMuPDF → Textract fallback), chunking, RAG store.

Vector layer is behind RagStore so ChromaDB can later be swapped for Pinecone/pgvector.
"""
import io
import os
import re
import uuid
import logging

log = logging.getLogger("doc_intel")

MIN_TEXT_CHARS = 200  # below this a PDF is considered scanned → OCR


# ── Extraction ────────────────────────────────────────────────────────────────
def _textract_images(images: list) -> str:
    """OCR page images with Amazon Textract (sync detect_document_text)."""
    import boto3
    client = boto3.client(
        "textract",
        region_name=os.environ.get("AWS_REGION", "us-east-1"),
        aws_access_key_id=os.environ.get("AWS_ACCESS_KEY_ID"),
        aws_secret_access_key=os.environ.get("AWS_SECRET_ACCESS_KEY"),
    )
    parts = []
    for img in images[:15]:
        try:
            resp = client.detect_document_text(Document={"Bytes": img})
            parts.append("\n".join(b.get("Text", "") for b in resp.get("Blocks", []) if b.get("BlockType") == "LINE"))
        except Exception as e:
            log.warning("textract page failed: %s", e)
    return "\n\n".join(p for p in parts if p)


def extract_any(data: bytes, filename: str) -> tuple:
    """Extract text from PDF/DOCX/TXT/CSV/image. Returns (text, method)."""
    name = (filename or "").lower()
    if name.endswith((".txt", ".csv", ".md")):
        return data.decode("utf-8", errors="replace")[:400000], "text"
    if name.endswith((".docx", ".doc")):
        try:
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(parts)[:400000], "docx"
        except Exception as e:
            raise ValueError(f"Could not read Word file: {e}")
    if name.endswith((".png", ".jpg", ".jpeg", ".webp")):
        txt = _textract_images([data]) if os.environ.get("AWS_ACCESS_KEY_ID") else ""
        return txt[:400000], "textract-image"
    # default: treat as PDF
    import fitz
    try:
        pdf = fitz.open(stream=data, filetype="pdf")
    except Exception:
        raise ValueError("Unsupported or corrupted file")
    text = "\n\n".join(pg.get_text() for pg in pdf)
    if len(text.strip()) >= MIN_TEXT_CHARS:
        return text[:400000], "pymupdf"
    # scanned PDF → render pages → Textract
    if os.environ.get("AWS_ACCESS_KEY_ID"):
        images = [pg.get_pixmap(dpi=150).tobytes("png") for pg in list(pdf)[:15]]
        ocr = _textract_images(images)
        if ocr.strip():
            return ocr[:400000], "textract"
    return text[:400000], "pymupdf-sparse"


# ── Chunking ──────────────────────────────────────────────────────────────────
def chunk_text(text: str, size: int = 1200, overlap: int = 150) -> list:
    text = re.sub(r"\n{3,}", "\n\n", text)
    chunks, i = [], 0
    while i < len(text):
        chunk = text[i:i + size]
        if chunk.strip():
            chunks.append(chunk.strip())
        i += size - overlap
        if len(chunks) >= 300:
            break
    return chunks


# ── RAG store (ChromaDB with keyword fallback) ────────────────────────────────
class RagStore:
    def __init__(self, path: str = "/app/backend/chroma_db"):
        self._col = None
        self._mem = []  # fallback keyword store: list of dicts
        try:
            import chromadb
            self._client = chromadb.PersistentClient(path=path)
            self._col = self._client.get_or_create_collection("bachein_docs")
        except Exception as e:
            log.warning("chromadb unavailable, keyword fallback: %s", e)

    def add(self, user_id: str, conv_id: str, document_id: str, file_name: str, chunks: list):
        metas = [{
            "user_id": user_id, "conversation_id": conv_id, "document_id": document_id,
            "file_name": file_name, "chunk_id": i, "page_number": -1,
        } for i in range(len(chunks))]
        ids = [f"{document_id}-{i}" for i in range(len(chunks))]
        if self._col is not None:
            try:
                self._col.add(documents=chunks, metadatas=metas, ids=ids)
                return
            except Exception as e:
                log.warning("chroma add failed: %s", e)
        for c, m in zip(chunks, metas):
            self._mem.append({"text": c, **m})

    def query(self, user_id: str, conv_id: str, q: str, k: int = 6) -> list:
        """Returns [{text, file_name, chunk_id}]"""
        if self._col is not None:
            try:
                r = self._col.query(
                    query_texts=[q], n_results=k,
                    where={"$and": [{"user_id": user_id}, {"conversation_id": conv_id}]},
                )
                docs = (r.get("documents") or [[]])[0]
                metas = (r.get("metadatas") or [[]])[0]
                return [{"text": d, "file_name": m.get("file_name", ""), "chunk_id": m.get("chunk_id", 0)} for d, m in zip(docs, metas)]
            except Exception as e:
                log.warning("chroma query failed: %s", e)
        # keyword fallback
        words = set(re.findall(r"[a-z]{3,}", q.lower()))
        scored = []
        for item in self._mem:
            if item["user_id"] != user_id or item["conversation_id"] != conv_id:
                continue
            score = sum(1 for w in words if w in item["text"].lower())
            if score:
                scored.append((score, item))
        scored.sort(key=lambda x: -x[0])
        return [{"text": it["text"], "file_name": it["file_name"], "chunk_id": it["chunk_id"]} for _, it in scored[:k]]


rag = RagStore()


# ── Exports (DOCX / HTML from plain text) ────────────────────────────────────
def text_to_docx(title: str, body: str) -> bytes:
    import docx
    d = docx.Document()
    d.add_heading(title, level=1)
    for para in body.split("\n\n"):
        if para.strip():
            d.add_paragraph(para.strip())
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def text_to_html(title: str, body: str) -> bytes:
    import html as _html
    paras = "".join(f"<p>{_html.escape(p.strip())}</p>\n" for p in body.split("\n\n") if p.strip())
    doc = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        f"<title>{_html.escape(title)}</title>"
        "<style>body{font-family:Georgia,serif;max-width:760px;margin:40px auto;line-height:1.7;color:#1a1915;padding:0 20px}"
        "h1{font-size:26px;border-bottom:2px solid #e8e5dc;padding-bottom:10px}</style></head>"
        f"<body><h1>{_html.escape(title)}</h1>\n{paras}</body></html>"
    )
    return doc.encode()

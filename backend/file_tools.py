"""File conversion utilities. Pure-python where possible."""
import io
import os
import tempfile
from typing import Tuple
from PIL import Image
from pypdf import PdfReader, PdfWriter
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage
from reportlab.lib.units import inch


SUPPORTED_INPUTS = ["pdf", "docx", "xlsx", "pptx", "txt", "md", "jpg", "jpeg", "png", "webp"]
SUPPORTED_OUTPUTS = ["pdf", "txt", "docx", "jpg", "png", "webp"]


def detect_format(filename: str) -> str:
    name = (filename or "").lower()
    for ext in SUPPORTED_INPUTS:
        if name.endswith("." + ext):
            return ext
    return ""


def _read_pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)
    except Exception:
        return ""


def _read_docx_text(data: bytes) -> str:
    bio = io.BytesIO(data)
    doc = DocxDocument(bio)
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _read_xlsx_text(data: bytes) -> str:
    bio = io.BytesIO(data)
    wb = load_workbook(bio, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            parts.append(" | ".join(str(c) if c is not None else "" for c in row))
    return "\n".join(parts)


def _read_pptx_text(data: bytes) -> str:
    bio = io.BytesIO(data)
    pres = Presentation(bio)
    parts = []
    for i, slide in enumerate(pres.slides, 1):
        parts.append(f"# Slide {i}")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
    return "\n\n".join(parts)


def _text_to_pdf_bytes(title: str, body: str) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()
    story = [Paragraph(f"<b>{_escape(title)}</b>", styles["Title"]), Spacer(1, 0.2 * inch)]
    for para in body.split("\n\n"):
        if para.strip():
            story.append(Paragraph(_escape(para).replace("\n", "<br/>"), styles["BodyText"]))
            story.append(Spacer(1, 0.12 * inch))
    doc.build(story)
    return buf.getvalue()


def _text_to_docx_bytes(body: str) -> bytes:
    docx = DocxDocument()
    for para in body.split("\n\n"):
        docx.add_paragraph(para)
    bio = io.BytesIO()
    docx.save(bio)
    return bio.getvalue()


def _image_bytes_to_pdf(data: bytes) -> bytes:
    img = Image.open(io.BytesIO(data))
    if img.mode != "RGB":
        img = img.convert("RGB")
    bio = io.BytesIO()
    img.save(bio, format="PDF")
    return bio.getvalue()


def _pdf_to_docx_bytes(data: bytes) -> bytes:
    """Use pdf2docx for real layout conversion."""
    from pdf2docx import Converter
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as inp:
        inp.write(data); pdf_path = inp.name
    docx_path = pdf_path.replace(".pdf", ".docx")
    try:
        cv = Converter(pdf_path)
        cv.convert(docx_path)
        cv.close()
        with open(docx_path, "rb") as f:
            return f.read()
    finally:
        for p in (pdf_path, docx_path):
            try: os.unlink(p)
            except Exception: pass


def convert(data: bytes, src_ext: str, target_ext: str, title: str = "Document") -> Tuple[bytes, str, str]:
    """Returns (bytes, content_type, filename)."""
    src = src_ext.lower().lstrip(".")
    tgt = target_ext.lower().lstrip(".")
    if src not in SUPPORTED_INPUTS:
        raise ValueError(f"Unsupported input format: .{src}")
    if tgt not in SUPPORTED_OUTPUTS:
        raise ValueError(f"Unsupported output format: .{tgt}")

    # Extract text representation
    text = ""
    if src == "pdf":
        text = _read_pdf_text(data)
    elif src == "docx":
        text = _read_docx_text(data)
    elif src == "xlsx":
        text = _read_xlsx_text(data)
    elif src == "pptx":
        text = _read_pptx_text(data)
    elif src in ("txt", "md"):
        text = data.decode("utf-8", errors="ignore")

    # Special direct conversions
    if src == "pdf" and tgt == "docx":
        return _pdf_to_docx_bytes(data), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"{title}.docx"

    # Image inputs/outputs
    if src in ("jpg", "jpeg", "png", "webp"):
        img = Image.open(io.BytesIO(data))
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
        if tgt == "pdf":
            return _image_bytes_to_pdf(data), "application/pdf", f"{title}.pdf"
        if tgt in ("jpg", "png", "webp"):
            bio = io.BytesIO()
            fmt = "JPEG" if tgt == "jpg" else tgt.upper()
            img.save(bio, format=fmt, quality=85)
            mime = {"jpg": "image/jpeg", "png": "image/png", "webp": "image/webp"}[tgt]
            return bio.getvalue(), mime, f"{title}.{tgt}"
        raise ValueError(f"Cannot convert image to .{tgt}")

    # Non-image targets
    if tgt == "pdf":
        return _text_to_pdf_bytes(title, text), "application/pdf", f"{title}.pdf"
    if tgt == "docx":
        return _text_to_docx_bytes(text), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"{title}.docx"
    if tgt == "txt":
        return text.encode("utf-8"), "text/plain", f"{title}.txt"

    raise ValueError(f"Cannot convert .{src} to .{tgt}")


def compress_image(data: bytes, quality: int = 60, max_dimension: int = 1600) -> Tuple[bytes, str]:
    img = Image.open(io.BytesIO(data))
    if img.mode in ("RGBA", "P"):
        img = img.convert("RGB")
    if max(img.size) > max_dimension:
        ratio = max_dimension / max(img.size)
        img = img.resize((int(img.size[0] * ratio), int(img.size[1] * ratio)))
    bio = io.BytesIO()
    img.save(bio, format="JPEG", quality=max(20, min(95, quality)), optimize=True)
    return bio.getvalue(), "image/jpeg"


def compress_pdf(data: bytes) -> bytes:
    reader = PdfReader(io.BytesIO(data))
    writer = PdfWriter()
    for page in reader.pages:
        try:
            page.compress_content_streams()  # may be slow on big files
        except Exception:
            pass
        writer.add_page(page)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def _escape(s: str) -> str:
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
